"""The deck as an editable PowerPoint - ten slides: native shapes, connectors, tables and pictures - no slide images.

    uv run --no-project --with python-pptx python3 docs/deck/build_pptx.py   ->   docs/contract-intelligence.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "contract-intelligence.pptx"
FONT = "Roboto"

TEAL, TEAL2, BG, PAPER, LINE = "00695C", "E0F2F1", "F4F6F6", "FFFFFF", "E3E7E7"
INK, MUTED, RED, RED2, AMBER, AMBER2, GREEN, GREEN2, GREY = "212121", "6B6B6B", "D32F2F", "FDECEA", "ED6C02", "FFF3E0", "2E7D32", "E8F5E9", "455A64"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- primitives
def _para(tf, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, first=False, italic=False, mono=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Menlo" if mono else FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    return p


def text(slide, x, y, w, h, lines, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, mono=False):
    """lines: str or list of (text, size, bold, color) tuples - one paragraph each."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    items = [(lines, size, bold, color)] if isinstance(lines, str) else lines
    for i, it in enumerate(items):
        t, s, b, c = (it + (size, bold, color))[:4] if isinstance(it, tuple) else (it, size, bold, color)
        _para(tf, t, s, b, c, align, first=(i == 0), mono=mono)
    return tb


def box(slide, x, y, w, h, title="", sub="", fill=PAPER, line=LINE, color=INK, size=12, subsize=10, subcolor=MUTED,
        bold=True, align=PP_ALIGN.CENTER, radius=0.12, line_w=1.0, anchor=MSO_ANCHOR.MIDDLE):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = min(0.5, radius / min(w, h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    if title:
        _para(tf, title, size, bold, color, align, first=True)
    if sub:
        _para(tf, sub, subsize, False, subcolor, align, first=not title)
    return shp


def pill(slide, x, y, label, on=False, w=None, size=10, color=None, fill=None):
    w = w or (0.34 + 0.1 * len(label))
    return box(slide, x, y, w, 0.3, label, fill=fill or (TEAL if on else TEAL2), line=None,
               color=color or ("FFFFFF" if on else TEAL), size=size, radius=0.15)


def arrow(slide, x1, y1, x2, y2, color="90A4AE", width=1.5, head=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def bar(slide, x, y, w, h, frac, color, bg):
    box(slide, x, y, w, h, fill=bg, line=None, radius=0.06)
    if frac > 0:
        box(slide, x, y, w * frac, h, fill=color, line=None, radius=0.06)


def table(slide, x, y, w, rows, widths, size=11, header=False, row_h=0.32):
    tbl = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PAPER)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            bold = header and r == 0
            col = INK if c == 0 or bold else MUTED
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = cell_text
            run.font.name, run.font.size, run.font.bold, run.font.color.rgb = FONT, Pt(size), bold or c == 0, rgb(col)
    tbl.first_row = header
    return tbl


def picture(slide, name, x, y, w, h):
    return slide.shapes.add_picture(str(HERE / "img" / name), Inches(x), Inches(y), Inches(w), Inches(h))


def new_slide(n, kicker, title):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    text(s, 0.6, 0.35, 12, 0.3, kicker.upper(), size=11, bold=True, color=TEAL)
    text(s, 0.6, 0.62, 12.1, 0.8, title, size=27, bold=True)
    text(s, 0.6, 7.05, 6, 0.3, "Contract Intelligence · Riverty case study", size=9, color=MUTED)
    text(s, 11.7, 7.05, 1, 0.3, f"{n} / 10", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return s




def poly(slide, points, color="90A4AE", width=1.5, head=True, dash=False):
    """An orthogonal connector drawn as straight segments; the arrowhead sits on the last one."""
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        arrow(slide, x1, y1, x2, y2, color=color, width=width, head=head and (x2, y2) == points[-1], dash=dash)



# ---------------------------------------------------------------- slides
# 1 title
s = new_slide(1, "Riverty · Software Engineer (RAG) case study", "Contract Intelligence")
text(s, 0.6, 1.9, 10, 1.2, [("Which contracts lack a clause the guideline requires?", 22, False, INK), ("Which still name the company by its old name?", 22, False, INK)])
text(s, 0.6, 3.15, 11, 0.5, "A knowledge-graph pipeline that finds, a strong model that verifies, a lawyer who decides — on the PDF itself.", size=14, color=MUTED)
xx = 0.6
for label, on in [("SharePoint · upload", False), ("Read", False), ("Neo4j graph", True), ("Check", False), ("Decide", False), ("Contract storage", False)]:
    w = 0.55 + 0.11 * len(label)
    box(s, xx, 4.1, w, 0.42, label, fill=TEAL if on else PAPER, line=None if on else LINE, color="FFFFFF" if on else INK, size=12, radius=0.1)
    xx += w + 0.15
    if label != "Contract storage":
        text(s, xx - 0.13, 4.13, 0.2, 0.4, "→", size=14, color=MUTED)
        xx += 0.2
text(s, 0.6, 5.6, 10, 0.4, "Gaurav Kulkarni · September 2026 · 10 minutes, then 5 minutes live", size=13, color=MUTED)

# 2 problem
s = new_slide(2, "The problem", "Open · search · verify — by hand, 15,000 times")
for i, (t, sub) in enumerate([("SharePoint", "most contracts, PDF"), ("Scans", "no text layer"), ("Handwritten JPEGs", "really old contracts"), ("Contract storage", "REST · store a copy only")]):
    box(s, 0.6, 1.7 + i * 0.95, 3.0, 0.8, t, sub, align=PP_ALIGN.LEFT, size=13)
box(s, 3.9, 1.7, 4.3, 1.2, "Which contracts do NOT contain passage X?", "QUESTION 1", size=16, subcolor=TEAL, align=PP_ALIGN.LEFT, line=TEAL, line_w=2)
box(s, 8.4, 1.7, 4.3, 1.2, "Which contracts still name the OLD company?", "QUESTION 2", size=16, subcolor=RED, align=PP_ALIGN.LEFT, line=RED, line_w=2)
box(s, 3.9, 3.15, 8.8, 2.3, fill=PAPER)
text(s, 4.05, 3.22, 5, 0.3, "Today, per contract and question", size=10, color=MUTED)
for i, t in enumerate(["open the PDF", "search, in two languages", "verify yourself"]):
    box(s, 4.3 + i * 2.85, 3.75, 2.3, 0.6, t, size=12)
    if i < 2:
        arrow(s, 6.6 + i * 2.85, 4.05, 7.15 + i * 2.85, 4.05)
poly(s, [(11.4, 4.05), (12.1, 4.05), (12.1, 3.5), (5.45, 3.5), (5.45, 3.75)], color=TEAL)
text(s, 10.9, 3.18, 1.5, 0.3, "× 15,000", size=11, bold=True, color=TEAL, align=PP_ALIGN.RIGHT)
text(s, 4.3, 4.6, 8, 0.3, "≈ 15 min · scans 20–40 · handwriting 30–60 or skipped", size=10, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 3.9, 5.6, 8.8, 0.7, "“Agents like Claude often fail due to the sheer size and complexity of the environments” — so this is not an agent. It is a pipeline with model nodes.", size=11, color=MUTED)

# 3 estimation
s = new_slide(3, "The estimation", "How much work is it — and what does it cost?")
tiles = [("15,000", "active contracts", "1,800 merchants × ~3 docs, vendors, mandates, NDAs"), ("40,000", "documents in the archive", "incl. arvato · infoscore · AfterPay entities"),
         ("12–20", "new contracts per day", "20–30 % of the stock turns over a year"), ("1–3", "sweeps a year", "rebrand 2022–23 · the bank 2026 · DORA · a guideline change")]
for i, (n, t, sub) in enumerate(tiles):
    x = 0.6 + i * 3.1
    box(s, x, 1.6, 2.95, 1.35, fill=PAPER)
    text(s, x + 0.12, 1.65, 2.7, 0.5, n, size=30, bold=True, color=TEAL)
    text(s, x + 0.12, 2.15, 2.7, 0.3, t, size=12, bold=True)
    text(s, x + 0.12, 2.42, 2.7, 0.5, sub, size=9, color=MUTED)
box(s, 0.6, 3.2, 7.0, 2.85, fill=PAPER)
text(s, 0.75, 3.28, 6.5, 0.3, "Rename sweep over 15,000 contracts", size=13, bold=True)
text(s, 0.75, 3.6, 3, 0.25, "manual · 15 min each", size=10, color=MUTED)
bar(s, 0.75, 3.88, 6.6, 0.28, 1.0, RED, RED2)
text(s, 0.75, 4.2, 6.6, 0.3, "3,750 h ≈ 2.3 FTE-years ≈ €340k        5 people · 5 months", size=11, bold=True)
text(s, 0.75, 4.75, 3.5, 0.25, "with the tool · 3 min per finding, spot checks", size=10, color=MUTED)
bar(s, 0.75, 5.03, 6.6, 0.28, 0.07, TEAL, TEAL2)
text(s, 0.75, 5.35, 6.6, 0.3, "250 h lawyer time + €900 model calls        1–2 weeks", size=11, bold=True)
box(s, 7.8, 3.2, 4.9, 2.85, fill=PAPER)
text(s, 7.95, 3.28, 4.6, 0.3, "Daily flow · 15 contracts a day", size=13, bold=True)
text(s, 7.95, 3.6, 3, 0.25, "manual · 45–60 min each", size=10, color=MUTED)
bar(s, 7.95, 3.88, 4.6, 0.28, 1.0, RED, RED2)
text(s, 7.95, 4.2, 4.6, 0.3, "1.6 FTE ≈ €240k a year", size=11, bold=True)
text(s, 7.95, 4.75, 3, 0.25, "with the tool", size=10, color=MUTED)
bar(s, 7.95, 5.03, 4.6, 0.28, 0.12, TEAL, TEAL2)
text(s, 7.95, 5.35, 4.6, 0.3, "0.2 FTE + €250 model + €3k infrastructure", size=11, bold=True)
text(s, 0.6, 6.25, 12.1, 0.6, "Gemini 3.1 Pro $2 / $12 per 1M tokens → ≈ $0.06–0.10 per contract. The tool costs ~1 % of the lawyer time it replaces; the time that remains is the part only a lawyer can do. Estimates: docs/cost-benefit.md.", size=10, color=MUTED)

# 4 thought process + versions
s = new_slide(4, "Thought process", "Retrieval finds what is there. It cannot prove what is not.")
box(s, 0.6, 1.6, 5.9, 2.75, fill=PAPER)
pill(s, 0.8, 1.72, "search", fill="ECEFF1", color=GREY)
for i, t in enumerate(["chunks", "similarity", "top-k"]):
    box(s, 0.9 + i * 1.85, 2.15, 1.55, 0.5, t)
    if i < 2:
        arrow(s, 2.45 + i * 1.85, 2.4, 2.75 + i * 1.85, 2.4)
text(s, 0.8, 2.85, 5.5, 0.3, "“nothing similar above 0.6” … is that absence?  →  a guess with a score", size=10, color=RED, align=PP_ALIGN.CENTER)
text(s, 0.8, 3.2, 5.5, 0.3, "“Arvato Systems” ≈ “Arvato Payment Solutions” to a vector  →  identity is not similarity", size=10, color=RED, align=PP_ALIGN.CENTER)
text(s, 0.8, 3.65, 5.5, 0.5, "an agent that opens and searches for itself fails in a big, messy environment — the statement says so", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
box(s, 6.8, 1.6, 5.9, 2.75, fill=PAPER, line=TEAL, line_w=2)
pill(s, 7.0, 1.72, "structure")
for i, (t, on) in enumerate([("clauses, typed", False), ("guideline REQUIRES", False), ("absence = lookup", True)]):
    box(s, 7.05 + i * 1.9, 2.15, 1.75, 0.5, t, fill=TEAL if on else TEAL2, line=None if on else TEAL, color="FFFFFF" if on else INK, size=10)
    if i < 2:
        arrow(s, 8.8 + i * 1.9, 2.4, 8.95 + i * 1.9, 2.4, color=TEAL)
for i, (t, on) in enumerate([("register (GLEIF)", False), ("RENAMED_TO", False), ("identity = edge", True)]):
    box(s, 7.05 + i * 1.9, 2.85, 1.75, 0.5, t, fill=TEAL if on else TEAL2, line=None if on else TEAL, color="FFFFFF" if on else INK, size=10)
    if i < 2:
        arrow(s, 8.8 + i * 1.9, 3.1, 8.95 + i * 1.9, 3.1, color=TEAL)
text(s, 7.0, 3.55, 5.5, 0.6, "rules FIND (recall)  →  a model VERIFIES (precision)  →  a person DECIDES", size=11, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
text(s, 0.6, 4.55, 6, 0.3, "Four versions in ten days", size=13, bold=True)
arrow(s, 0.9, 5.1, 12.4, 5.1, color=TEAL, width=3, head=False)
cards = [("26–27 Aug · v1 workbench", "question cards, matrix, review queue, chat — every feature defensible, the sum not what a legal team reaches for: cut", False),
         ("29 Aug · v2 file converter", "drop in → one result line → findings on the PDF → corrected copy: lawyers trust what they see on the page", False),
         ("5 Sep · v3 knowledge graph", "Neo4j as the RAG store beside Postgres; register from GLEIF; the models read the graph’s selection", False),
         ("5 Sep · v4 graph only", "Neo4j the sole database; scoped-then-full verification; three steps + a Technik page: one system, precision restored", True)]
for i, (t, body, on) in enumerate(cards):
    x = 0.6 + i * 3.1
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.32), Inches(4.97), Inches(0.26), Inches(0.26))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(TEAL if on else PAPER)
    c.line.color.rgb = rgb(TEAL)
    c.line.width = Pt(2.5)
    box(s, x, 5.4, 2.95, 1.45, fill=PAPER, line=TEAL if on else LINE, line_w=2 if on else 1)
    text(s, x + 0.1, 5.45, 2.75, 1.35, [(t, 10, True, TEAL), (body, 9, False, INK)])

# 5 workflow
s = new_slide(5, "The workflow", "One contract, end to end — 14 stages, three phases")
box(s, 0.6, 3.05, 1.35, 0.8, "SharePoint", "upload · delta sync", size=11, subsize=9)
arrow(s, 1.95, 3.45, 2.3, 3.45)
box(s, 2.3, 1.6, 3.1, 4.3, fill=PAPER)
text(s, 2.3, 1.68, 3.1, 0.3, "READ", size=12, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
for i, t in enumerate(["load · text layer or page image", "ocr · Tesseract → Gemini vision", "segment · clauses by structure", "classify · keywords + Flash", "entities · register + GLEIF", "screen · injection guard", "embed · vector + full-text"]):
    box(s, 2.5, 2.05 + i * 0.53, 2.7, 0.42, t, fill=TEAL2, line=TEAL, size=10, bold=False)
arrow(s, 5.4, 3.45, 5.8, 3.45)
box(s, 5.8, 2.55, 2.1, 1.8, "Neo4j", "contracts · clauses · names\nguideline · register · decisions\nvector + full-text index", fill=TEAL, line=None, color="FFFFFF", subcolor="FFFFFF", size=16, subsize=9)
arrow(s, 7.9, 3.45, 8.3, 3.45)
box(s, 8.3, 1.6, 3.0, 4.3, fill=PAPER)
text(s, 8.3, 1.68, 3.0, 0.3, "CHECK · LangGraph", size=12, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
for i, t in enumerate(["rules · guideline gaps, old names", "cross_check · Pro, two reads", "place · box on the page", "draft · name or clause", "policy · carry-over, spot checks"]):
    box(s, 8.5, 2.05 + i * 0.53, 2.6, 0.42, t, fill=TEAL2, line=TEAL, size=10, bold=False)
arrow(s, 11.3, 3.45, 11.65, 3.45)
box(s, 11.65, 2.55, 1.15, 1.8, fill=PAPER)
text(s, 11.65, 2.62, 1.15, 0.3, "DECIDE", size=12, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
box(s, 11.75, 3.0, 0.95, 0.42, "decide", fill=TEAL2, line=TEAL, size=10, bold=False)
box(s, 11.75, 3.55, 0.95, 0.42, "correct", fill=TEAL2, line=TEAL, size=10, bold=False)
arrow(s, 12.22, 4.35, 12.22, 4.85)
box(s, 11.55, 4.85, 1.35, 0.65, "contract storage", size=10)
poly(s, [(9.8, 5.9), (9.8, 6.3), (6.85, 6.3), (6.85, 4.35)], color=TEAL, dash=True)
text(s, 7.0, 6.35, 3, 0.3, "decisions become precedents", size=9, color=MUTED)
text(s, 0.6, 6.65, 12, 0.3, "So funktioniert es and Technik render this same stage list from GET /api/pipeline — the explanation cannot drift from the code", size=10, color=MUTED, align=PP_ALIGN.CENTER)

# 6 graph — a grid with orthogonal edges only
s = new_slide(6, "The graph", "The database is the knowledge graph")
box(s, 0.6, 1.6, 8.1, 5.2, fill=PAPER)
NW, NH = 1.9, 0.62
X = {"A": 0.9, "B": 3.45, "C": 6.0}
Y = {1: 1.95, 2: 3.15, 3: 4.35}


def gnode(col, row, t, sub="", fill=PAPER, line=LINE, color=INK, lw=1.0):
    return box(s, X[col], Y[row], NW, NH, t, sub, fill=fill, line=line, color=color, size=11, subsize=8, radius=0.31, line_w=lw)


def mid(col, row):
    return X[col] + NW / 2, Y[row] + NH / 2


gnode("B", 2, "Contract", fill=TEAL, line=None, color="FFFFFF")
gnode("B", 1, "ContractType", "merchant agreement", fill=TEAL2, line=TEAL)
gnode("C", 1, "ClauseType", "liability_cap", fill=TEAL2, line=TEAL)
gnode("C", 2, "ClauseType", "anti_corruption", fill=TEAL2, line=RED, lw=2.5)
gnode("B", 3, "Clause", "§ 4 · vector · text")
gnode("A", 2, "Entity", "Arvato Payment Sol. GmbH")
gnode("A", 3, "Entity", "Riverty GmbH · LEI", fill=TEAL2, line=TEAL)
# OF_TYPE: Contract up to ContractType
arrow(s, mid("B", 2)[0], Y[2], mid("B", 1)[0], Y[1] + NH)
text(s, mid("B", 2)[0] + 0.05, Y[1] + NH + 0.12, 1.2, 0.3, "OF_TYPE", size=9, color=MUTED)
# REQUIRES: ContractType right to liability_cap, and a bus down to anti_corruption
arrow(s, X["B"] + NW, mid("B", 1)[1], X["C"], mid("C", 1)[1])
text(s, X["B"] + NW + 0.05, mid("B", 1)[1] - 0.32, 1.0, 0.3, "REQUIRES", size=9, color=MUTED)
poly(s, [(X["B"] + NW + 0.25, mid("B", 1)[1]), (X["B"] + NW + 0.25, mid("C", 2)[1]), (X["C"], mid("C", 2)[1])])
# MENTIONS: Contract left to Entity; RENAMED_TO: Entity down
arrow(s, X["B"], mid("B", 2)[1], X["A"] + NW, mid("A", 2)[1])
text(s, X["A"] + NW + 0.05, mid("A", 2)[1] - 0.32, 1.2, 0.3, "MENTIONS", size=9, color=MUTED)
arrow(s, mid("A", 2)[0], Y[2] + NH, mid("A", 3)[0], Y[3], color=TEAL)
text(s, mid("A", 2)[0] + 0.05, Y[2] + NH + 0.12, 1.3, 0.3, "RENAMED_TO", size=9, bold=True, color=TEAL)
# HAS_CLAUSE: Contract down to Clause; IS_A: Clause right, up the far column, into liability_cap
arrow(s, mid("B", 2)[0], Y[2] + NH, mid("B", 3)[0], Y[3])
text(s, mid("B", 2)[0] + 0.05, Y[2] + NH + 0.12, 1.3, 0.3, "HAS_CLAUSE", size=9, color=MUTED)
bus = X["C"] + NW + 0.25
poly(s, [(X["B"] + NW, mid("B", 3)[1]), (bus, mid("B", 3)[1]), (bus, mid("C", 1)[1]), (X["C"] + NW, mid("C", 1)[1])], color=TEAL)
text(s, bus + 0.05, mid("C", 1)[1] + 0.2, 0.6, 0.3, "IS_A", size=9, bold=True, color=TEAL)
arrow(s, bus, mid("C", 2)[1], X["C"] + NW, mid("C", 2)[1], color=RED, dash=True)
text(s, X["C"] - 0.2, Y[2] + NH + 0.05, 2.4, 0.3, "no IS_A path → the gap", size=10, bold=True, color=RED, align=PP_ALIGN.CENTER)
text(s, 0.75, 5.35, 7.8, 0.3, "(Decision)-[:ON]->(Contract), (Decision)-[:ABOUT]->(ClauseType): the team’s decisions live in the same graph as precedents", size=9, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.75, 5.85, 7.8, 0.3, "MATCH (c)-[:OF_TYPE]->()-[:REQUIRES]->(t) WHERE NOT EXISTS { (c)-[:HAS_CLAUSE]->()-[:IS_A]->(t) }", size=9, color=TEAL, align=PP_ALIGN.CENTER, mono=True)
text(s, 0.75, 6.15, 7.8, 0.3, "“which contracts lack a required clause” — one pattern, no search, no model", size=9, color=MUTED, align=PP_ALIGN.CENTER)
cards = [("database-first", "Seeded from files and registers, not from a model: the guideline (Legal’s file), the taxonomy, GLEIF — the public LEI register records Arvato Payment Solutions GmbH as the previous name of Riverty GmbH — and the team’s decisions."),
         ("retrieval inside the graph", "Every clause carries its embedding (768 d, cosine, Cypher SEARCH) and its text (Lucene). Hybrid = reciprocal rank fusion of both; scoped by graph filters."),
         ("why not …", "Postgres + pgvector (v1–v3): the questions became joins over label tables. Cosmos DB Gremlin: no native vectors. AGE on Azure PostgreSQL: the Microsoft runner-up — swaps in for two modules.")]
for i, (tag, body) in enumerate(cards):
    y = 1.6 + i * 1.75
    box(s, 8.9, y, 3.8, 1.6, fill=PAPER)
    pill(s, 9.05, y + 0.12, tag)
    text(s, 9.02, y + 0.45, 3.55, 1.1, body, size=9.5)

# 7 check + decide
s = new_slide(7, "Check · Decide", "Rules for recall, the model for precision")
box(s, 0.6, 1.6, 6.2, 5.2, fill=PAPER)
box(s, 0.8, 1.8, 5.8, 0.75, "candidates from the rules", "required type without a certain clause · active old-name mention — recall 0.95, precision 0.43", fill=TEAL2, line=TEAL, size=11, subsize=8.5)
arrow(s, 3.7, 2.55, 3.7, 2.8)
box(s, 0.8, 2.8, 5.8, 1.0, "read 1 · what the graph selects", "outline (every clause, type, page) + the 6 most relevant clauses + every clause of the type;\nGemini Pro, thinking high → confirmed | refuted | partial + page + quote + reason", size=11, subsize=8.5)
arrow(s, 3.7, 3.8, 3.7, 4.05)
box(s, 2.5, 4.05, 2.4, 0.45, "still “missing”?", line=AMBER, line_w=2, size=11, radius=0.22)
arrow(s, 4.9, 4.28, 6.0, 4.28, color=TEAL)
text(s, 5.0, 3.98, 1.6, 0.3, "refuted → present", size=8.5, bold=True, color=TEAL)
arrow(s, 3.7, 4.5, 3.7, 4.75)
box(s, 0.8, 4.75, 5.8, 0.7, "read 2 · the whole contract", "skipped when read 1 already covered ≥ 90 % · the cheap error (a false alarm) is caught here", fill=TEAL2, line=TEAL, size=11, subsize=8.5)
arrow(s, 3.7, 5.45, 3.7, 5.7)
box(s, 1.6, 5.7, 4.2, 0.5, "confirmed → the page · precision 1.00 · recall 0.91 (CUAD, 22 absences)", fill=TEAL, line=None, color="FFFFFF", size=9.5)
text(s, 0.8, 6.3, 5.8, 0.4, "only the verifier may remove a candidate · Pro for verify and handwriting, Flash for labels and drafts, Flash-Lite for screening · Foundry behind the same switch", size=8, color=MUTED, align=PP_ALIGN.CENTER)
box(s, 7.0, 1.6, 5.7, 3.35, fill=PAPER)
text(s, 7.15, 1.68, 5.4, 0.3, "A person decides — the review rate falls with evidence", size=12, bold=True)
ox, oy, gw, gh = 7.75, 2.15, 4.7, 1.9
arrow(s, ox, oy + gh, ox + gw, oy + gh, color=LINE, head=False)
arrow(s, ox, oy, ox, oy + gh, color=LINE, head=False)
for lbl, frac in [("100 %", 1.0), ("20 %", 0.2), ("10 %", 0.1)]:
    text(s, ox - 0.65, oy + gh - frac * gh - 0.13, 0.6, 0.3, lbl, size=8, color=MUTED, align=PP_ALIGN.RIGHT)
pts = [(0, 1.0), (0.18, 1.0), (0.18, 0.2), (0.6, 0.2), (0.6, 0.1), (0.85, 0.1), (0.85, 1.0), (1.0, 1.0)]
for (a, b), (c, d) in zip(pts, pts[1:]):
    arrow(s, ox + a * gw, oy + gh - b * gh, ox + c * gw, oy + gh - d * gh, color=TEAL, width=2.5, head=False)
for lbl, fx in [("8 agreeing · ≥ 95 %", 0.18), ("24 in a row", 0.6), ("a dismissal → reset", 0.85)]:
    text(s, ox + fx * gw - 0.9, oy + gh + 0.05, 1.8, 0.3, lbl, size=8, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 7.15, 4.45, 5.4, 0.45, "deterministic spot checks (SHA-256), at least one per contract; only cross-checked findings are automated; decisions carry over per file and become precedents", size=8.5, color=MUTED)
box(s, 7.0, 5.15, 2.75, 1.65, fill=PAPER)
text(s, 7.12, 5.2, 2.5, 0.3, "a false alarm, reviewed", size=9, color=MUTED)
text(s, 7.12, 5.45, 2.5, 0.4, "€4.50", size=18, bold=True, color=GREEN)
text(s, 7.12, 5.85, 2.5, 0.4, "3 minutes on the page", size=9, color=MUTED)
box(s, 9.95, 5.15, 2.75, 1.65, fill=PAPER)
text(s, 10.07, 5.2, 2.5, 0.3, "a false edit, filed", size=9, color=MUTED)
text(s, 10.07, 5.45, 2.5, 0.4, "€1k–100k", size=18, bold=True, color=RED)
text(s, 10.07, 5.85, 2.5, 0.6, "rework, dispute — a document with legal effect", size=9, color=MUTED)

# 8 UI
s = new_slide(8, "The UI", "Three steps — the PDF is the interface")
cols = [("1 · Hochladen", "drop one or many · one result line each · a filter for “which contracts lack X”", "01-home.jpg", 3.4),
        ("2 · Prüfen", "markers on the page · one finding at a time · Übernehmen / Nicht zutreffend", "02-contract.jpg", 4.9),
        ("3 · Herunterladen", "what was applied · the copy · preview · filing", "02b-download.jpg", 3.4)]
x = 0.6
for tag, sub, im, w in cols:
    pill(s, x, 1.55, tag)
    text(s, x, 1.85, w, 0.5, sub, size=9, color=MUTED)
    picture(s, im, x, 2.35, w, 2.8)
    x += w + 0.2
whys = [("Why one finding at a time.", "The only clicks are decisions; a decision moves on; the marker on the page is verified with a glance — a list would send the lawyer back into the PDF."),
        ("Why no numbers for lawyers.", "No confidences, no traces on these pages — quote, reason, marker, and whether it was cross-checked. German first, English one click."),
        ("Why a Technik page.", "Engineers get the same 14 stages with module, model, thresholds, the graph’s live counts and a trace of one contract — generated, so it cannot drift.")]
for i, (h, b) in enumerate(whys):
    x = 0.6 + i * 4.1
    box(s, x, 5.4, 3.9, 1.4, fill=PAPER)
    text(s, x + 0.12, 5.47, 3.65, 1.3, [(h, 10, True, INK), (b, 9.5, False, INK)])

# 9 infrastructure + alignment
s = new_slide(9, "Infrastructure · the problem statement", "Azure — and the statement, checked")
box(s, 0.6, 1.6, 6.9, 5.2, fill=PAPER)
env = box(s, 2.5, 1.85, 4.85, 4.75, fill=PAPER, line=TEAL, radius=0.15)
env.line.dash_style = 4
text(s, 2.5, 1.9, 4.85, 0.3, "Azure · Container Apps · Terraform", size=10, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
box(s, 0.8, 2.55, 1.5, 0.75, "SharePoint", "Graph delta", size=10, subsize=8)
arrow(s, 2.3, 2.92, 2.75, 2.92)
box(s, 2.75, 2.4, 1.9, 1.0, "API", "FastAPI · LangGraph\ninternal ingress", fill=TEAL2, line=TEAL, size=11, subsize=8)
box(s, 2.75, 3.85, 1.9, 0.8, "Web", "React · Entra ID sign-in", fill=TEAL2, line=TEAL, size=11, subsize=8)
arrow(s, 3.7, 3.85, 3.7, 3.4)
box(s, 5.2, 2.4, 1.9, 1.0, "Neo4j", "AuraDB (Marketplace)\nor the container", fill=TEAL, line=None, color="FFFFFF", subcolor="FFFFFF", size=12, subsize=8)
arrow(s, 4.65, 2.9, 5.2, 2.9)
for i, t in enumerate(["Blob · working copies", "Key Vault", "Log Analytics", "Document Intelligence"]):
    box(s, 5.2, 3.65 + i * 0.62, 1.9, 0.5, t, size=9.5)
poly(s, [(4.95, 3.4), (4.95, 3.65 + 3 * 0.62 + 0.25)], color="B0BEC5", width=1, head=False)
arrow(s, 4.65, 3.15, 4.95, 3.15, color="B0BEC5", width=1, head=False)
for i in range(4):
    arrow(s, 4.95, 3.65 + i * 0.62 + 0.25, 5.2, 3.65 + i * 0.62 + 0.25, color="B0BEC5", width=1)
box(s, 0.8, 3.85, 1.5, 0.75, "Gemini", "Vertex · EU", size=10, subsize=8)
box(s, 0.8, 4.75, 1.5, 0.75, "Foundry", "GPT-5 · Data Zone", size=10, subsize=8)
box(s, 0.8, 5.65, 1.5, 0.75, "Contract storage", "Idempotency-Key", size=9, subsize=8)
poly(s, [(2.1, 3.3), (2.1, 5.9)], color="B0BEC5", width=1, head=False)
arrow(s, 2.75, 3.3, 2.1, 3.3, color="B0BEC5", width=1, head=False)
for yy in (4.22, 5.12, 6.02):
    arrow(s, 2.1, yy, 2.3, yy, color="B0BEC5", width=1)
text(s, 2.55, 6.3, 4.8, 0.3, "documents never leave the tenant; only the checked contract’s pages go to the model", size=8, color=MUTED, align=PP_ALIGN.CENTER)
box(s, 7.7, 1.6, 5.0, 5.2, fill=PAPER)
text(s, 7.85, 1.68, 4.7, 0.3, "Against the problem statement", size=12, bold=True)
rows = [("✓", "basic pipeline · front-end · end-to-end workflow"), ("✓", "contracts without a passage · old company name"), ("✓", "scans and handwritten JPEGs · Tesseract → vision"),
        ("✓", "best-of-breed vs Microsoft: one switch each"), ("✓", "Azure · Foundry/Gemini · containers · FastAPI · Terraform · React/TS · LangChain/LangGraph"),
        ("◐", "SharePoint: Graph delta sync written, not run against a tenant"), ("◐", "contract storage: idempotent filing, mocked API"), ("◐", "guidelines · best interests: guideline file, drafts, no playbook yet")]
for i, (m, a) in enumerate(rows):
    y = 2.1 + i * 0.5
    pill(s, 7.9, y, m, w=0.35, fill=GREEN2 if m == "✓" else AMBER2, color=GREEN if m == "✓" else AMBER)
    text(s, 8.35, y - 0.03, 4.25, 0.5, a, size=9.5)
text(s, 7.85, 6.15, 4.7, 0.6, "Switches: LLM gemini | foundry · OCR tesseract | document_intelligence · source folder | sharepoint · database Neo4j | AGE + pgvector on Azure PostgreSQL", size=8, color=MUTED)

# 10 adoption + demo
s = new_slide(10, "Adoption · live demo", "Adoption — then five minutes live")
steps = [("1 · Pilot on the next sweep", "the bank rename, two lawyers, measured against the manual sample", False), ("2 · Legal owns the rules", "the guideline file, the register, the clause types — no engineer needed", False),
         ("3 · Decide on the page", "every decision teaches: precedents, the review rate as the trust meter", False), ("4 · Routine", "sweeps on events, the daily flow, spot checks, the corrected copy filed", True)]
for i, (t, sub, on) in enumerate(steps):
    x = 0.6 + i * 3.1
    box(s, x, 1.6, 2.8, 1.05, t, sub, fill=TEAL if on else TEAL2, line=None if on else TEAL, color="FFFFFF" if on else INK, subcolor="FFFFFF" if on else MUTED, size=11, subsize=8.5)
    if i < 3:
        arrow(s, x + 2.8, 2.12, x + 3.1, 2.12)
text(s, 0.6, 2.8, 12, 0.5, "trust: the original is never touched, every decision is a named person’s, the audit log answers who accepted what · learning: So funktioniert es and Technik, generated from the code · measure: minutes per contract, review rate per class, precision/recall against the answer key", size=8.5, color=MUTED)
box(s, 0.6, 3.5, 7.0, 3.3, fill=PAPER)
text(s, 0.75, 3.58, 3, 0.3, "Six beats", size=12, bold=True)
beats = ["drop a contract → “Wird gelesen … Wird geprüft …” → the result line", "open it: the box on the page, the quote, the model’s reason, Übernehmen / Nicht zutreffend",
         "Weiter zum Download: the corrected copy, the preview, filing under an idempotency key", "the start page filter: which contracts lack a limitation of liability?",
         "Technik: the stages with their models, the graph’s live counts, the gap query, a trace", "the Neo4j browser: the contract as a graph"]
for i, b in enumerate(beats):
    y = 3.95 + i * 0.44
    pill(s, 0.8, y, str(i + 1), w=0.35, size=9)
    text(s, 1.25, y - 0.03, 6.2, 0.42, b, size=10)
box(s, 7.8, 3.5, 4.9, 3.3, fill=PAPER)
text(s, 7.95, 3.58, 3, 0.3, "Next", size=12, bold=True)
nxt = ["SharePoint against the real tenant", "the contract storage’s real API — an adapter", "Entra ID sign-in, structured logging, metrics", "a playbook of preferred positions for the drafts", "parallel ingestion behind a queue for sweeps", "a pairwise “compare two contracts” view"]
text(s, 7.95, 3.95, 4.6, 2.2, [("· " + n, 10, False, INK) for n in nxt])
text(s, 7.95, 6.2, 4.6, 0.5, "docs/architecture-decisions.md · docs/cost-benefit.md · docs/ps-coverage.md", size=8, color=MUTED)

prs.save(OUT)
print(f"wrote {OUT} — {len(prs.slides._sldIdLst)} slides")
