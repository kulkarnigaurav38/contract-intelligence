"""Three slides, few words, editable shapes: the problem, the pipeline, the tech."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

TEAL = RGBColor(0x00, 0x69, 0x5C)
TEAL_LIGHT = RGBColor(0xE0, 0xF2, 0xF0)
INK = RGBColor(0x1F, 0x29, 0x29)
MUTED = RGBColor(0x6B, 0x77, 0x77)
LINE = RGBColor(0xC9, 0xD3, 0xD2)
RED = RGBColor(0xC6, 0x28, 0x28)
RED_LIGHT = RGBColor(0xFD, 0xEC, 0xEA)
AMBER = RGBColor(0xB2, 0x6A, 0x00)
AMBER_LIGHT = RGBColor(0xFF, 0xF3, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF7, 0xF9, 0xF9)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = s.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def _no_shadow(sh):
    from pptx.oxml.ns import qn
    spPr = sh._element.spPr
    for e in spPr.findall(qn("a:effectLst")):
        spPr.remove(e)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    style = sh._element.find(qn("p:style"))  # the theme style is what renderers turn into a drop shadow
    if style is not None:
        sh._element.remove(style)


def box(slide, x, y, w, h, s, fill=WHITE, stroke=TEAL, color=INK, size=13, bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, sub=None, sub_color=MUTED):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = stroke
    sh.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.18
    _no_shadow(sh)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = s
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(max(8, size - 4))
        r2.font.color.rgb = sub_color
        r2.font.name = "Consolas"
    return sh


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def header(slide, kicker, title, n):
    text(slide, 0.6, 0.35, 8, 0.3, kicker.upper(), size=11, bold=True, color=TEAL)
    text(slide, 0.6, 0.6, 11.5, 0.9, title, size=30, bold=True, color=INK, font="Georgia")
    text(slide, 11.9, 0.4, 0.9, 0.3, f"{n} / 3", size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    text(slide, 0.6, 7.0, 8, 0.3, "Contract Intelligence · Riverty case study", size=9, color=MUTED)


# ---------------------------------------------------------------- 1 · the problem
s = prs.slides.add_slide(BLANK)
header(s, "The problem", "Repetitive contract review, done by hand.", 1)

# sources
box(s, 0.7, 2.0, 2.3, 1.0, "SharePoint", sub="PDF", fill=PAPER, stroke=LINE, size=15, bold=True)
box(s, 0.7, 3.3, 2.3, 1.0, "Old scans", sub="JPEG · handwritten", fill=PAPER, stroke=LINE, size=15, bold=True)
box(s, 0.7, 4.6, 2.3, 1.0, "Contract storage", sub="REST · store-only copy", fill=PAPER, stroke=LINE, size=15, bold=True)

# the manual loop
loop = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.2), Inches(2.0), Inches(3.6), Inches(3.6))
loop.fill.solid(); loop.fill.fore_color.rgb = RED_LIGHT; loop.line.color.rgb = RED; loop.line.width = Pt(1.5); _no_shadow(loop)
loop.text_frame.text = ""
box(s, 5.1, 2.25, 1.8, 0.6, "open", fill=WHITE, stroke=RED, color=RED, size=14, bold=True)
box(s, 5.1, 3.5, 1.8, 0.6, "search", fill=WHITE, stroke=RED, color=RED, size=14, bold=True)
box(s, 5.1, 4.75, 1.8, 0.6, "verify", fill=WHITE, stroke=RED, color=RED, size=14, bold=True)
arrow(s, 6.0, 2.85, 6.0, 3.5, color=RED)
arrow(s, 6.0, 4.1, 6.0, 4.75, color=RED)
text(s, 4.2, 5.7, 3.6, 0.4, "every contract · every question", size=11, color=RED, align=PP_ALIGN.CENTER)
arrow(s, 3.0, 2.5, 4.3, 3.2)
arrow(s, 3.0, 3.8, 4.2, 3.8)

# the two questions
box(s, 8.9, 2.3, 3.8, 1.2, "Which contracts lack a clause?", fill=TEAL_LIGHT, stroke=TEAL, size=16, bold=True)
box(s, 8.9, 3.9, 3.8, 1.2, "Where is the old company name?", fill=TEAL_LIGHT, stroke=TEAL, size=16, bold=True)
text(s, 8.9, 5.3, 3.8, 0.5, "arvato Financial Solutions → Riverty GmbH", size=11, color=MUTED, align=PP_ALIGN.CENTER)
arrow(s, 7.8, 3.8, 8.9, 3.8)

# ---------------------------------------------------------------- 2 · the pipeline
s = prs.slides.add_slide(BLANK)
header(s, "The pipeline", "Drop a contract. Get the findings on the page.", 2)

phases = [
    ("READ", 0.6, TEAL, TEAL_LIGHT, [
        ("Read file", "PyMuPDF · text or image"), ("OCR", "Tesseract → Gemini Vision"), ("Split", "into clauses"),
        ("Label", "12 types · rules + Flash"), ("Names", "old-name register"), ("Screen", "prompt injection"),
        ("Index", "pgvector · full text"),
    ], ""),
    ("CHECK", 4.9, AMBER, AMBER_LIGHT, [
        ("Rules", "guideline × clauses"), ("Cross-check", "Gemini Pro reads it all"),
        ("Place", "box on the page"), ("Draft", "new name · clause"), ("Policy", "earlier decisions"),
    ], "only the cross-check may remove a finding\nquote · page · reasoning"),
    ("DECIDE", 9.2, RED, RED_LIGHT, [
        ("Decide", "on the PDF · accept / no"), ("Corrected copy", "original untouched"), ("Store", "REST · idempotent"),
    ], "decisions are remembered:\nsame file → carried over\nsame kind → fewer checks"),
]
W, ROW, GAP = 3.55, 0.5, 0.13
for name, x, color, fill, steps, note in phases:
    band = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.75), Inches(W), Inches(4.95))
    band.fill.solid(); band.fill.fore_color.rgb = fill; band.line.color.rgb = LINE; band.line.width = Pt(0.75)
    band.adjustments[0] = 0.04; _no_shadow(band)
    text(s, x + 0.15, 1.85, W - 0.3, 0.35, name, size=11, bold=True, color=color)
    y = 2.25
    for i, (label, sub) in enumerate(steps):
        sh = box(s, x + 0.2, y, W - 0.4, ROW, label, fill=WHITE, stroke=color, size=12, bold=True)
        tf = sh.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        r = tf.paragraphs[0].add_run(); r.text = "   " + sub; r.font.size = Pt(9); r.font.color.rgb = MUTED; r.font.name = "Consolas"; r.font.bold = False
        if i < len(steps) - 1:
            arrow(s, x + W / 2, y + ROW, x + W / 2, y + ROW + GAP, color=MUTED, width=1)
        y += ROW + GAP
    if note:
        text(s, x + 0.2, 6.7 - 0.22 * (note.count("\n") + 1) - 0.1, W - 0.4, 0.22 * (note.count("\n") + 1) + 0.1, note, size=10, color=INK)
arrow(s, 4.15, 3.6, 4.9, 3.6, color=INK, width=2.25)
arrow(s, 8.45, 3.6, 9.2, 3.6, color=INK, width=2.25)

# ---------------------------------------------------------------- 3 · the tech
s = prs.slides.add_slide(BLANK)
header(s, "The tech", "Best-of-breed today, Microsoft behind one switch.", 3)

layers = [
    ("Front-end", ["React + TypeScript", "MUI"], None),
    ("API", ["Python · FastAPI", "LangChain + LangGraph", "PyMuPDF · Tesseract"], None),
    ("Models", ["Gemini 3.x  Pro · Flash · Flash-Lite · embeddings"], "Azure Foundry · Document Intelligence"),
    ("Data", ["PostgreSQL + pgvector", "full-text index"], None),
    ("Sources", ["SharePoint (Graph delta)", "Contract storage REST API"], None),
    ("Infra", ["Docker · Azure Container Apps", "Terraform · Key Vault · Entra ID"], None),
]
y = 1.75
for name, items, alt in layers:
    box(s, 0.7, y, 1.7, 0.7, name, fill=TEAL, stroke=TEAL, color=WHITE, size=13, bold=True)
    x = 2.6
    for it in items:
        w = 0.12 * len(it) + 0.5
        box(s, x, y, w, 0.7, it, fill=WHITE, stroke=LINE, size=12)
        x += w + 0.15
    if alt:
        box(s, 9.0, y, 3.7, 0.7, alt, sub="LLM_PROVIDER=foundry · OCR_PROVIDER=document_intelligence", fill=PAPER, stroke=LINE, color=MUTED, size=11)
        arrow(s, 8.55, y + 0.35, 9.0, y + 0.35, color=MUTED, width=1)
    y += 0.82
text(s, 0.7, 6.75, 12, 0.3, "one routing table: task → model tier → thinking level  ·  same code path for Gemini and Azure OpenAI", size=10, color=MUTED)

out = str(__import__("pathlib").Path(__file__).with_name("contract-intelligence.pptx"))
prs.save(out)
print("saved", out)
